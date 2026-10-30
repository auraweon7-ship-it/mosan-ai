# -*- coding: utf-8 -*-
"""엑셀 교육계획(52주 상세일정 시트)을 읽어 config.js의 주차 목록을 다시 씁니다.

사용법:
    pip install openpyxl
    python tools/excel_to_config.py [엑셀 파일 경로]

엑셀 경로를 생략하면 data/schedule.xlsx 를 읽습니다.
config.js 안의 /* WEEKS:BEGIN */ … /* WEEKS:END */ 사이만 바꾸고, 나머지 내용은 그대로 둡니다.
"""
import datetime
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG = os.path.join(ROOT, "config.js")
SHEET = "52주 상세일정"
HEADER = "주차"
# 제목이 이 형식에 맞으면 '발표' 꼬리표를 붙입니다. ('발표자료 제작'처럼 만드는 주차는 제외)
PRESENT_RE = re.compile(r"발표(?!자료)|시연")


def read_weeks(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[1]
    rows = list(ws.iter_rows(values_only=True))
    head_idx = next(i for i, r in enumerate(rows) if r and r[0] == HEADER)
    weeks = []
    for r in rows[head_idx + 1:]:
        if not r or not isinstance(r[0], (int, float)):
            continue
        no, date, _dow, time, mode, stage, title, details, practice, outcome = (list(r) + [None] * 10)[:10]
        if isinstance(date, datetime.datetime):
            date = date.strftime("%Y-%m-%d")
        topics = [re.sub(r"^[•·\-\s]+", "", t).strip() for t in str(details or "").splitlines()]
        tag = "발표" if PRESENT_RE.search(str(title)) else ""
        weeks.append({
            "week": int(no),
            "title": str(title or "").strip(),
            "tag": tag,
            "stage": str(stage or "").strip(),
            "date": date,
            "time": str(time or "").replace("~", " – ").strip(),
            "location": str(mode or "").strip(),
            "topics": [t for t in topics if t],
            "practice": str(practice or "").strip(),
            "outcome": str(outcome or "").strip(),
            "materials": [],
            "videos": [],
        })
    return weeks


def to_js(weeks):
    lines = []
    for w in weeks:
        lines.append("      " + json.dumps(w, ensure_ascii=False) + ",")
    return "\n".join(lines)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "data", "schedule.xlsx")
    weeks = read_weeks(src)
    with open(CONFIG, encoding="utf-8") as f:
        text = f.read()
    pattern = re.compile(r"(/\* WEEKS:BEGIN \*/\n).*?([ \t]*/\* WEEKS:END \*/)", re.S)
    if not pattern.search(text):
        sys.exit("config.js에서 WEEKS:BEGIN / WEEKS:END 표시를 찾지 못했습니다.")
    text = pattern.sub(lambda m: m.group(1) + to_js(weeks) + "\n" + m.group(2), text)
    with open(CONFIG, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)
    print("%d개 주차를 config.js에 반영했습니다 (%s ~ %s)." % (len(weeks), weeks[0]["date"], weeks[-1]["date"]))


if __name__ == "__main__":
    main()
